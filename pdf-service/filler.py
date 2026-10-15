"""Word 範本填入資料 → 轉 PDF → 合併（Cloud Run 與本機共用）"""
import copy, io, os, re, subprocess, tempfile, shutil, hashlib, threading, zipfile
import docx
from docx.shared import Cm
from docx.oxml.ns import qn
from docx_tools import iter_paragraphs, set_text

TAG = re.compile(r'\{\{([^}]+)\}\}')
DATE_KEYS = ['到職日', '最後工作日', '生日', '簽署日']

def expand_dates(v):
    out = dict(v)
    for k in DATE_KEYS:
        val = v.get(k)
        if isinstance(val, str) and re.match(r'^\d{4}-\d{2}-\d{2}$', val):
            y, m, d = [int(x) for x in val.split('-')]
            out[k + '_民國年'] = str(y - 1911); out[k + '_月'] = str(m); out[k + '_日'] = str(d)
            out[k + '_民國'] = f'{y - 1911}年{m}月{d}日'
    return out

def render_text(text, v):
    def rep(m):
        key = m.group(1).strip()
        if key.startswith('chk:'):
            k, want = key[4:].split('=', 1)
            have = v.get(k)
            on = (want in have) if isinstance(have, list) else (str(have) == want)
            return '■' if on else '□'
        if key == '簽名': return m.group(0)
        return str(v.get(key, ''))
    return TAG.sub(rep, text)

def _float_picture(run, height_cm, raise_cm):
    """把剛插入的行內圖片改成「浮動、不換行」：不撐開行高、不推動版面（LibreOffice 對 line 定位不準，改用段落定位）"""
    from docx.oxml import parse_xml
    inline = run._r.find('.//' + qn('wp:inline'))
    if inline is None: return
    ext = inline.find(qn('wp:extent'))
    cx, cy = ext.get('cx'), ext.get('cy')
    docpr = inline.find(qn('wp:docPr')); graphic = inline.find(qn('a:graphic'))
    off = int(raise_cm * 360000)   # 從段落頂端往下的距離（負值＝往上）
    anchor = parse_xml(
        '<wp:anchor xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
        'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        'distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="251659264" behindDoc="0" locked="0" layoutInCell="1" allowOverlap="1">'
        '<wp:simplePos x="0" y="0"/>'
        '<wp:positionH relativeFrom="character"><wp:posOffset>0</wp:posOffset></wp:positionH>'
        f'<wp:positionV relativeFrom="paragraph"><wp:posOffset>{off}</wp:posOffset></wp:positionV>'
        f'<wp:extent cx="{cx}" cy="{cy}"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/>'
        '</wp:anchor>')
    anchor.append(docpr)
    anchor.append(parse_xml('<wp:cNvGraphicFramePr xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"/>'))
    anchor.append(graphic)
    inline.getparent().replace(inline, anchor)

def trim_sig(png):
    """裁掉簽名畫布四周的空白，只留筆跡（簽名才會貼在正確位置）"""
    try:
        from PIL import Image, ImageOps
        im = Image.open(io.BytesIO(png)).convert('RGBA')
        a = im.getchannel('A')
        if a.getextrema()[0] == 255:   # 沒有透明度 → 用「非白色」判斷筆跡
            a = ImageOps.invert(im.convert('L')).point(lambda v: 255 if v > 40 else 0)
        box = a.getbbox()
        if not box: return png
        pad = 6; w, h = im.size
        im = im.crop((max(0, box[0]-pad), max(0, box[1]-pad), min(w, box[2]+pad), min(h, box[3]+pad)))
        out = io.BytesIO(); im.save(out, 'PNG'); return out.getvalue()
    except Exception:
        return png

def put_signature(p, sig_png, height_cm=1.2, max_w_cm=4.0):
    before, after = p.text.split('{{簽名}}', 1)
    set_text(p, before)
    rpr = p.runs[0]._r.find(qn('w:rPr')) if p.runs else None
    from PIL import Image
    w, h = Image.open(io.BytesIO(sig_png)).size
    if h * 4 <= w or w < 4: w_cm = max_w_cm; h_cm = max(0.3, w_cm * h / max(w, 1))   # 很扁的簽名
    else: h_cm = height_cm; w_cm = min(max_w_cm, h_cm * w / h)
    if w_cm == max_w_cm: h_cm = w_cm * h / w
    r = p.add_run()
    r.add_picture(io.BytesIO(sig_png), width=Cm(w_cm), height=Cm(h_cm))
    _float_picture(r, h_cm, 1.1 - h_cm)   # 簽名底部約落在該行文字基線
    sz = rpr.find(qn('w:sz')) if rpr is not None else None
    ch_cm = (int(sz.get(qn('w:val'))) / 2 if sz is not None else 12) * 0.0353
    tail = '\u3000' * (int(w_cm / ch_cm) + 1) + (after or '\u2060')   # 留出簽名寬度，後面文字不會被蓋住；結尾加「零寬連接字」避免靠右對齊時行尾空白被忽略而超出邊界
    r2 = p.add_run(tail)
    if rpr is not None: r2._r.insert(0, copy.deepcopy(rpr))

def fill_docx(template_bytes, values, sig_png):
    sig_png = trim_sig(sig_png)
    d = docx.Document(io.BytesIO(template_bytes))
    for p in list(iter_paragraphs(d)):
        t = p.text
        if '{{' not in t: continue
        new = render_text(t, values)
        if new != t: set_text(p, new)
        if '{{簽名}}' in p.text: put_signature(p, sig_png)
    out = io.BytesIO(); d.save(out); return out.getvalue()

# LibreOffice：共用設定檔目錄（只在第一次建立，省下每次 10 秒以上），同一時間只跑一個
LO_PROFILE = 'file:///tmp/lo_profile'
_LO_LOCK = threading.Lock()

def _soffice(tmp, srcs, timeout=150):
    with _LO_LOCK:
        return subprocess.run(['soffice', '--headless', '--norestore', '--nologo', '--nodefault', f'-env:UserInstallation={LO_PROFILE}',
                               '--convert-to', 'pdf', '--outdir', tmp, *srcs], capture_output=True, text=True, timeout=timeout)

def warm_up():
    """容器啟動時先跑一次空白轉檔，建立 LibreOffice 設定檔並載入字型"""
    tmp = tempfile.mkdtemp()
    try:
        d = docx.Document(); d.add_paragraph('warm up 暖機'); src = os.path.join(tmp, 'w.docx'); d.save(src)
        _soffice(tmp, [src], timeout=170)
    except Exception:
        pass
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

def _pages(pdf):
    r = subprocess.run(['pdfinfo', pdf], capture_output=True, text=True)
    m = re.search(r'Pages:\s+(\d+)', r.stdout)
    return int(m.group(1)) if m else 0

def compact_docx(b, f=0.85):
    """把行距、段距、列高等比例縮小（不改字級）：讓只多出幾行的單頁表單擠回一頁"""
    zin = zipfile.ZipFile(io.BytesIO(b)); out = io.BytesIO()
    zout = zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED)
    def sc(m): return f'{m.group(1)}="{max(0, int(int(m.group(2)) * f))}"'
    def sp(m):
        t = m.group(0)
        t = re.sub(r'(w:(?:before|after))="(\d+)"', lambda x: f'{x.group(1)}="{int(int(x.group(2)) * f * 0.6)}"', t)
        if 'w:lineRule="auto"' in t or 'lineRule' not in t:
            t = re.sub(r'w:line="(\d+)"', lambda x: f'w:line="{min(int(x.group(1)), 240) if int(x.group(1)) > 100 else x.group(1)}"', t)
        else:
            t = re.sub(r'(w:line)="(\d+)"', sc, t)
        return t
    for it in zin.infolist():
        data = zin.read(it.filename)
        if it.filename == 'word/document.xml':
            x = data.decode('utf8')
            x = re.sub(r'<w:spacing [^>]*w:(?:line|before|after)=[^>]*/>', sp, x)
            x = re.sub(r'(<w:trHeight w:val)="(\d+)"', sc, x)
            data = x.encode('utf8')
        zout.writestr(it, data)
    zout.close(); return out.getvalue()

def _convert_all(tmp, templates, v, sig_png):
    """所有文件一次交給 LibreOffice 轉檔（只啟動一次）；回傳 [(name, pdf_path)]
    原本 2 頁的文件會再試一次「縮行距」版本，若能變成 1 頁就改用（Word 單頁表單在伺服器字型下常多出幾行）"""
    srcs, filled = [], []
    for i, (n, d) in enumerate(templates):
        src = os.path.join(tmp, f'{i:02d}.docx')
        fb = fill_docx(d, v, sig_png); filled.append(fb)
        open(src, 'wb').write(fb); srcs.append(src)
    r = _soffice(tmp, srcs)
    out = []
    for i, (n, _) in enumerate(templates):
        pdf = os.path.join(tmp, f'{i:02d}.pdf')
        if not os.path.exists(pdf):
            raise RuntimeError(f'轉檔失敗 {n}: rc={r.returncode} {r.stderr[-300:]}')
        out.append((n, pdf))
    two = [i for i, (_, pdf) in enumerate(out) if _pages(pdf) == 2]
    if two:
        csrc = []
        for i in two:
            src = os.path.join(tmp, f'{i:02d}c.docx'); open(src, 'wb').write(compact_docx(filled[i])); csrc.append(src)
        try:
            _soffice(tmp, csrc)
            for i in two:
                cp = os.path.join(tmp, f'{i:02d}c.pdf')
                if os.path.exists(cp) and _pages(cp) == 1: out[i] = (out[i][0], cp)
        except Exception:
            pass
    return out

def build_pdf(templates, values, sig_png):
    """templates: [(檔名, docx bytes)]；回傳 (合併後 PDF bytes, sha256)"""
    v = expand_dates(values)
    tmp = tempfile.mkdtemp()
    try:
        pdfs = [p for _, p in _convert_all(tmp, templates, v, sig_png)]
        merged = os.path.join(tmp, 'merged.pdf')
        if len(pdfs) == 1: shutil.copy(pdfs[0], merged)
        else: subprocess.run(['pdfunite', *pdfs, merged], check=True)
        b = open(merged, 'rb').read()
        return b, hashlib.sha256(b).hexdigest()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

def build_previews(templates, values, sig_png, dpi=72, quality=55, with_pdf=False):
    """每份文件各自轉成頁面 JPEG；回傳 [(name, [jpeg bytes...])]；with_pdf=True 時另回傳合併 PDF bytes"""
    v = expand_dates(values)
    tmp = tempfile.mkdtemp()
    try:
        out, pdfs = [], []
        for i, (n, pdf) in enumerate(_convert_all(tmp, templates, v, sig_png)):
            pdfs.append(pdf)
            prefix = os.path.join(tmp, f'p{i:02d}')
            subprocess.run(['pdftoppm', '-r', str(dpi), '-jpeg', '-jpegopt', f'quality={quality}', pdf, prefix], check=True, timeout=120)
            pages = sorted(f for f in os.listdir(tmp) if f.startswith(f'p{i:02d}') and f.endswith('.jpg'))
            out.append((n, [open(os.path.join(tmp, f), 'rb').read() for f in pages]))
        if not with_pdf: return out
        merged = os.path.join(tmp, 'merged.pdf')
        if len(pdfs) == 1: shutil.copy(pdfs[0], merged)
        else: subprocess.run(['pdfunite', *pdfs, merged], check=True)
        return out, open(merged, 'rb').read()
    finally:
        shutil.rmtree(tmp, ignore_errors=True)
